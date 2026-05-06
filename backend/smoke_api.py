#!/usr/bin/env python3
"""
Basic smoke-test script for the Inkforge API backend.
Run with: python test_api.py
"""
import requests
import json
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

BASE = "http://localhost:8000"

def test(name, fn):
    try:
        result = fn()
        print(f"  ✅ {name}: {result}")
    except Exception as e:
        print(f"  ❌ {name}: {e}")

print("\n🔍 Inkforge API — Smoke Tests")
print("=" * 50)

# Health
print("\n[Health]")
test("GET /api/health", lambda: requests.get(f"{BASE}/api/health", timeout=3).json()["status"])

# Templates
print("\n[Templates]")
test("GET /api/templates", lambda: f"{len(requests.get(f'{BASE}/api/templates', timeout=3).json())} templates")

# Genres
print("\n[Genres]")
test("GET /api/genres", lambda: f"{len(requests.get(f'{BASE}/api/genres', timeout=3).json())} genres")

# Market
print("\n[Market]")
for genre_id in ["literary_fiction", "thriller", "fantasy"]:
    test(f"GET /api/market/{genre_id}", lambda g=genre_id: requests.get(f"{BASE}/api/market/{g}", timeout=3).json()["name"])

# Query
print("\n[Query Generation]")
def test_manual_query():
    data = {"title": "Test Title", "author_name": "Test Author", "word_count": 80000}
    res = requests.post(f"{BASE}/api/query/manual", data={"data": json.dumps(data)})
    if res.headers.get("content-type") != "application/zip":
        raise Exception("Expected ZIP file")
    
    import zipfile
    import io
    from docx import Document
    
    with zipfile.ZipFile(io.BytesIO(res.content)) as z:
        for filename in z.namelist():
            if filename.endswith(".docx"):
                with z.open(filename) as f:
                    doc = Document(f)
                    text = " ".join([p.text for p in doc.paragraphs])
                    if "[Write your hook" in text:
                        raise Exception(f"Placeholder found in {filename}")
    return "Valid ZIP with no placeholders"

test("POST /api/query/manual (Empty fields no placeholders)", test_manual_query)

print("\n" + "=" * 50)
print("Smoke tests complete. Run 'uvicorn main:app --reload' in the backend folder first.\n")
